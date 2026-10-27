"""Rebuild the Polish DERA-HRI v3 Word questionnaire from the canonical CSV.

Usage: python tools/export_questionnaire_docx.py
Requires python-docx only for document export, not for research analysis.
"""
import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "instruments" / "dera_hri_v3_pl.csv"
DEST = ROOT / "instruments" / "dera_hri_v3_pl.docx"
SCALES = {
    "EAR-D": "Ucieleśniona reaktywność afektywna",
    "SEI-D": "Interpretacja semiotyczno-ucieleśniona",
    "SAP-D": "Predykcja społeczno-afektywna",
    "MP-D": "Przypisywanie stanów umysłowych",
    "RAR-D": "Relacyjna regulacja afektywna",
    "RM": "Mechanizmy regulacji (moduł dodatkowy)",
    "FII": "Przewidywany wpływ przyszłych interakcji (moduł dodatkowy)",
}
OPTIONS = (
    "Zdecydowanie się nie zgadzam",
    "Nie zgadzam się",
    "Raczej się nie zgadzam",
    "Ani się zgadzam, ani nie zgadzam",
    "Raczej się zgadzam",
    "Zgadzam się",
    "Zdecydowanie się zgadzam",
)
NAVY = RGBColor(16, 65, 98)


def main():
    with SOURCE.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream))
    if [r["item_id"] for r in rows] != [
        *(str(i) for i in range(1, 46)), "M1", "M2", "M3", "F1"
    ]:
        raise ValueError("Unexpected Polish item bank structure")
    document = Document()
    sec = document.sections[0]
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    base = document.styles["Normal"]
    base.font.name = "Aptos"
    base.font.size = Pt(9.8)
    base.paragraph_format.space_after = Pt(3)
    for style_name, size in (("Title", 22), ("Heading 1", 13)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = NAVY

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    head = header.add_run("DERA-HRI  |  WARIANT POLSKI  |  v3.0-PL")
    head.font.size = Pt(8)
    head.font.bold = True
    head.font.color.rgb = NAVY
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Tłumaczenie robocze · 27.10.2026 · Brak walidacji")
    run.font.size = Pt(7.5)

    document.add_heading("DERA-HRI v3.0", 0)
    p = document.add_paragraph(
        "Kwestionariusz doświadczeń afektywnych w interakcji człowiek–robot | wariant polski"
    )
    p.style = document.styles["Subtitle"]
    p = document.add_paragraph()
    note = p.add_run(
        "STATUS: TŁUMACZENIE ROBOCZE. Nie jest to zwalidowana adaptacja "
        "językowa ani narzędzie z zatwierdzonym kluczem punktowania."
    )
    note.bold = True
    note.font.size = Pt(8.5)
    note.font.color.rgb = RGBColor(135, 60, 34)
    document.add_heading("Instrukcja dla osoby badanej", 1)
    document.add_paragraph(
        "Przed chwilą zakończono interakcję z robotem. Poniższe stwierdzenia "
        "dotyczą wyłącznie Twoich odczuć i doświadczeń podczas tej konkretnej "
        "interakcji, a nie ogólnych opinii o robotach. Przy każdym stwierdzeniu "
        "zaznacz jedną odpowiedź. Nie ma odpowiedzi poprawnych ani błędnych. "
        "Jeżeli robot nie korzystał z kanału ekspresji opisanego w pytaniu, "
        "powiedz o tym prowadzącemu badanie zamiast wybierać odpowiedź neutralną."
    )
    document.add_paragraph(
        "Kod sesji: __________________  Data: __________  "
        "Robot / scenariusz: __________________"
    )
    document.add_heading("Skala odpowiedzi", 1)
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for i, label in enumerate(OPTIONS, 1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = label
    document.add_paragraph(
        "Każda pozycja używa tej samej skali 1–7. Jeśli pytanie nie dotyczy "
        "interakcji, zaznacz N/D i przekaż powód prowadzącemu."
    )

    prev_scale = None
    for item in rows:
        scale = item["scale"]
        if scale != prev_scale:
            document.add_heading(f"{scale}  |  {SCALES[scale]}", 1)
            if scale == "SEI-D":
                document.add_paragraph(
                    "Pytania 12–15 dotyczą dostępnych kanałów ruchu, głosu, "
                    "gestów i mimiki; pytanie 16 wymaga porównania kilku kanałów."
                )
            if scale in {"RM", "FII"}:
                document.add_paragraph(
                    "Moduł dodatkowy, oceniany osobno. "
                    "Nie wliczać do żadnej podskali podstawowej."
                )
            prev_scale = scale
        para = document.add_paragraph()
        para.paragraph_format.space_before = Pt(2)
        run = para.add_run(item["item_id"] + ".  ")
        run.bold = True
        run.font.color.rgb = NAVY
        para.add_run(item["text_pl"])
        answer = document.add_paragraph(
            "□ 1      □ 2      □ 3      □ 4      □ 5      □ 6      □ 7      □ N/D"
        )
        answer.paragraph_format.left_indent = Cm(1.1)
        answer.paragraph_format.space_after = Pt(1)
        for run in answer.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = NAVY

    document.add_heading("Informacja metodologiczna", 1)
    document.add_paragraph(
        "Wariant v3.0-PL opracowano jako tłumaczenie angielskiego banku "
        "DERA-HRI v3.0 (45 pozycji głównych i cztery dodatkowe). "
        "Pozostawiono oryginalne identyfikatory i niejednoznaczności "
        "sformułowań. Przed wnioskowaniem naukowym wymagane są niezależna "
        "weryfikacja językowa, wywiady poznawcze i ocena równoważności "
        "pomiarowej. Nie należy łączyć odpowiedzi z wariantów EN i PL ani "
        "wyliczać niezatwierdzonych wyników."
    )
    document.add_paragraph(
        "Repozytorium: https://github.com/MatPomGit/arc-dera-hri"
    )
    document.save(DEST)
    print(f"Created {DEST} ({len(rows)} items)")


if __name__ == "__main__":
    main()
